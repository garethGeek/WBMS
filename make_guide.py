from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'E:\PROJECTS\TANROADS\wbms\WBMS_Indicator_Profile_Guide.docx'
GREEN = '1D6F59'; NAVY = '10233F'; PALE = 'EAF5F0'; GRAY = '5F6B78'

def cell_style(cell, fill=None):
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(shd)
    mar = OxmlElement('w:tcMar')
    for side in ('top','start','bottom','end'):
        x = OxmlElement('w:' + side); x.set(qn('w:w'), '140'); x.set(qn('w:type'), 'dxa'); mar.append(x)
    cell._tc.get_or_add_tcPr().append(mar)

def bullet(doc, text): doc.add_paragraph(text, style='List Bullet')
def numbered(doc, text): doc.add_paragraph(text, style='List Number')
def note(doc, label, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; c = t.cell(0,0); cell_style(c, PALE)
    p = c.paragraphs[0]; r = p.add_run(label + '  '); r.bold = True; r.font.color.rgb = RGBColor.from_string(GREEN); p.add_run(text)

doc = Document(); s = doc.sections[0]; s.top_margin = Inches(.75); s.bottom_margin = Inches(.7); s.left_margin = Inches(.85); s.right_margin = Inches(.85)
n = doc.styles['Normal']; n.font.name = 'Aptos'; n.font.size = Pt(10.5); n.font.color.rgb = RGBColor.from_string(NAVY); n.paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 28, NAVY), ('Heading 1', 17, GREEN)]:
    st = doc.styles[name]; st.font.name = 'Aptos Display'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
header = s.header.paragraphs[0]; header.text = 'WBMS  |  TANROADS'; header.runs[0].font.size = Pt(8); header.runs[0].font.bold = True; header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
footer = s.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; footer.text = 'Indicator Profile Guide'; footer.runs[0].font.size = Pt(8); footer.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(20); r = p.add_run('WBMS TECHNICAL GUIDE'); r.bold = True; r.font.color.rgb = RGBColor.from_string(GREEN)
doc.add_paragraph('Indicator Profiles', style='Title'); doc.add_paragraph('How indicator profiles work with the TANROADS Weighbridge Management System')
note(doc, 'Key principle:', 'An indicator profile is a reusable decoder configuration. It tells WBMS how to recognize and interpret an indicator frame; it does not replace communication settings or repair hardware or network problems.')

doc.add_heading('1. What Is an Indicator Profile?', 1)
doc.add_paragraph('An indicator profile is a saved configuration describing how a particular weighing indicator sends its data. Different manufacturers, models, and installations may use different frame structures even when they measure the same weight.')
doc.add_paragraph('A profile can store:')
for x in ['Manufacturer and model information', 'Serial or Ethernet/TCP communication details used during testing', 'Start and end bytes, such as STX (02) and ETX (03)', 'The separator between multiple scale or axle values, such as SPACE, comma, or literal 049', 'Single-cell or multi-cell scale count', 'String length and characters removed from the beginning or end', 'Serial parameters such as baud rate, data bits, parity, stop bits, and flow control']: bullet(doc, x)

doc.add_heading('2. How It Works in WBMS', 1)
doc.add_paragraph('The workflow has two stages: the profile is tested and saved, then it is explicitly assigned to a registered device.')
for x in ['Open Indicator Profile Creator and enter connection details.', 'Connect and inspect ASCII Raw Data and HEX Raw Data.', 'Tune framing, separator, scale count, and trimming until Parsed Weight is correct.', 'Save the tested configuration as a reusable profile.', 'Open Device Registry and explicitly select the profile for the device.', 'Start live capture. WBMS applies the assigned profile while reading that device.']: numbered(doc, x)
note(doc, 'Assignment rule:', 'Saving a profile does not automatically change any registered device. It affects a device only after it is selected and the device configuration is saved.')

doc.add_heading('3. What Happens During Live Capture?', 1)
for x in ['Connect using the device registry connection settings.', 'Locate a complete frame using configured start and end bytes.', 'Remove framing bytes before weight extraction.', 'Apply beginning/end trimming and string-length limits.', 'Split multi-cell data using the configured separator and scale count.', 'Convert valid numeric values into kilograms.', 'Publish total weight and individual cell or axle values to the dashboard.']: numbered(doc, x)

doc.add_heading('4. Example Formats', 1)
t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
for i, h in enumerate(['Indicator output','Profile interpretation','WBMS result']):
    c=t.rows[0].cells[i]; c.text=h; cell_style(c,GREEN); c.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); c.paragraphs[0].runs[0].bold=True
for row in [('ST,GS,+001410 kg','Single numeric weight','1,410 kg'),('<STX>565 300 600 100<ETX>','Four values separated by spaces','1,565 kg total'),('<STX>566049300049601049100<ETX>','Four values separated by literal 049','1,567 kg total')]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v; cell_style(cells[i])

doc.add_heading('5. Main Function', 1); doc.add_paragraph('The profile is the translation layer between bytes sent by the indicator and standardized kilogram readings displayed by WBMS. It solves indicator-format differences without rewriting the capture engine for every model.')
doc.add_heading('6. Advantages', 1)
for x in ['Supports different manufacturers and models.', 'Allows testing before assignment to a live device.', 'Provides ASCII and hexadecimal diagnostics.', 'Supports single-cell and multi-cell readings.', 'Allows reusable configurations.', 'Makes separator, framing, and trimming adjustments testable.', 'Keeps built-in protocols separate from custom profiles.', 'Allows profiles to be edited as the actual frame becomes known.']: bullet(doc, x)
doc.add_heading('7. Disadvantages and Current Limitations', 1)
for x in ['Cannot correct an incorrect IP, TCP port, COM port, baud rate, or wiring problem.', 'A changed indicator frame may require profile editing.', 'An incorrect separator or scale count can produce an incorrect total.', 'ASCII and HEX panels show actual received bytes; they are not rewritten.', 'The current system focuses on framing, trimming, separators, and cell extraction, not every checksum or command-response protocol.', 'Weight token index is stored, but current multi-cell extraction primarily relies on separators and scale count.']: bullet(doc, x)

doc.add_heading('8. Built-In Protocol vs. Indicator Profile', 1)
t = doc.add_table(rows=1, cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['Area','Built-in protocol','Indicator profile']):
    c=t.rows[0].cells[i]; c.text=h; cell_style(c,NAVY); c.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255); c.paragraphs[0].runs[0].bold=True
for row in [('Configuration','Defined in WBMS code','Defined and saved by an authorized user'),('Assignment','Selected as device protocol','Explicitly assigned to a device'),('Flexibility','Stable for known formats','Tunable for an observed frame'),('Best use','Known manufacturer protocol','Custom or site-specific format')]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v; cell_style(cells[i])

doc.add_heading('9. Shared Understanding', 1); doc.add_paragraph('A saved indicator profile is a reusable decoder, not an automatic device assignment. It becomes active only when an operator selects it for a specific device and saves that device. Once active, it controls how WBMS frames, trims, separates, and translates indicator output into kilograms. If no profile is assigned, the device uses its selected built-in protocol.')
note(doc, 'Operational reminder:', 'Confirm the ASCII frame, HEX bytes, scale count, separator, and calculated kilogram result before assigning a profile to a production weighbridge device.')
doc.core_properties.title='WBMS Indicator Profiles Guide'; doc.core_properties.subject='Indicator profile behavior and WBMS integration'; doc.core_properties.author='TANROADS WBMS'; doc.save(OUT); print(OUT)
